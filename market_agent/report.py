"""Writes the formatted Excel report."""
import pandas as pd
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SIGNAL_FILLS = {
    "STRONG UP": PatternFill("solid", fgColor="00B050"),
    "UP": PatternFill("solid", fgColor="C6EFCE"),
    "NEUTRAL": PatternFill("solid", fgColor="EDEDED"),
    "DOWN": PatternFill("solid", fgColor="FFC7CE"),
    "STRONG DOWN": PatternFill("solid", fgColor="FF5B5B"),
}

PARAMETER_GUIDE = [
    # (Group, Parameter, What it tells you, Bullish reading, Bearish reading, Source)
    ("Trend", "50 / 200-day moving average (DMA)", "Medium & long-term trend", "Price above DMA; 50-DMA above 200-DMA (golden cross)", "Price below DMA; death cross", "Yahoo Finance prices"),
    ("Trend", "Supertrend (10, 3)", "Popular Kite/TradingView trend-following signal", "BUY (price above band)", "SELL", "Calculated"),
    ("Trend", "ADX (14) & +DI/-DI", "Strength of trend (above 25 = strong)", "ADX>25 with +DI > -DI", "ADX>25 with -DI > +DI", "Calculated"),
    ("Trend", "EMA 9 vs EMA 21", "Short-term trend", "EMA9 above EMA21", "EMA9 below EMA21", "Calculated"),
    ("Momentum", "RSI (14)", "Speed of price moves (0-100)", "Below 30 (oversold bounce) or 50-68 (healthy momentum)", "Above 75 (overbought)", "Calculated"),
    ("Momentum", "MACD (12,26,9)", "Momentum shift", "Bullish crossover / rising positive histogram", "Bearish crossover / falling negative histogram", "Calculated"),
    ("Momentum", "Stochastic (14,3)", "Close vs recent range", "%K crossing up from below 20", "%K crossing down from above 80", "Calculated"),
    ("Momentum", "Williams %R, CCI (20)", "Overbought / oversold", "Deeply negative, turning up", "Very high, turning down", "Calculated"),
    ("Momentum", "Returns 1d/5d/20d/60d/120d", "Recent performance (momentum effect)", "Positive, consistent", "Negative", "Calculated"),
    ("Momentum", "Relative strength vs NIFTY", "Is the stock beating the market?", "Outperforming", "Underperforming", "Calculated"),
    ("Volatility", "Bollinger Bands (20,2) %B & width", "Stretch & squeeze", "Below lower band; breakout after squeeze", "Above upper band (stretched)", "Calculated"),
    ("Volatility", "ATR (14)", "Typical daily move -> expected range & stop loss", "-", "-", "Calculated"),
    ("Price action", "20-day breakout / breakdown", "New highs or lows", "Close above 20-day high", "Close below 20-day low", "Calculated"),
    ("Price action", "Distance from 52-week high/low", "Where price sits in its yearly range", "Near 52-week high", "Near 52-week low with weak trend", "Calculated"),
    ("Price action", "Gap %, close position in day's range", "Buyer/seller control at close", "Close near day high", "Close near day low", "Calculated"),
    ("Volume", "Volume surge vs 20-day average", "Institutional activity", "Big volume on up day", "Big volume on down day", "Yahoo Finance"),
    ("Volume", "OBV (On-Balance Volume)", "Accumulation vs distribution", "Rising", "Falling", "Calculated"),
    ("Volume", "MFI (14)", "Volume-weighted RSI", "Below 20", "Above 80", "Calculated"),
    ("Volume", "Delivery % (NSE bhavcopy)", "Share of volume actually taken into demat - genuine buying", "High delivery (1.3x avg) on up day", "High delivery on down day", "NSE archives"),
    ("Fundamental", "P/E, P/B", "Valuation", "Reasonable", "P/E above 100", "Yahoo Finance"),
    ("Fundamental", "ROE, ROA, margins", "Business quality", "ROE above 18%", "ROE below 5%, losses", "Yahoo Finance"),
    ("Fundamental", "Debt / Equity", "Balance-sheet risk", "Below 0.5", "Above 2", "Yahoo Finance"),
    ("Fundamental", "Earnings & revenue growth (YoY)", "Growth", "Above +10%", "Below -10%", "Yahoo Finance"),
    ("Fundamental", "Promoter / institutional holding", "Ownership quality", "High", "Low", "Yahoo Finance"),
    ("Fundamental", "Analyst target & rating", "Street consensus", "Target 15%+ above price; BUY", "Price above target; SELL", "Yahoo Finance"),
    ("Event", "Upcoming quarterly results", "Earnings can move price 5-15% either way", "Flagged in 'Event Alert' column", "-", "Yahoo Finance"),
    ("Market", "NIFTY 50 trend & RSI", "Most stocks follow the index", "NIFTY above 50-DMA", "NIFTY below 50-DMA", "Yahoo Finance"),
    ("Market", "India VIX", "Fear gauge", "Low / falling", "Above 20 / rising sharply", "Yahoo Finance"),
    ("Model", "ML probability (Gradient Boosting)", "Learns from 3 yrs of history how all the above combined predicted the next 7 trading days", "Above 55%", "Below 45%", "Trained nightly"),
]

MAIN_COLS = {
    "symbol": "Symbol", "company": "Company", "industry": "Industry", "exchange": "Exchange",
    "close": "Close (Rs)", "ret_1d": "1D %", "prediction": "Prediction", "prob_up": "Prob. Up",
    "confidence": "Confidence", "expected_ret": "Expected Move %", "target_price": "Target (Rs)",
    "likely_low": "Likely Low (Rs)", "likely_high": "Likely High (Rs)", "stop_loss": "Stop Loss (Rs)",
    "ml_prob_up": "ML Prob. Up", "rule_score": "Rule Score", "technical_score": "Tech Score",
    "fundamental_score": "Fund. Score", "event_alert": "Event Alert", "reasons": "Key Reasons",
}
TECH_COLS = {
    "symbol": "Symbol", "close": "Close", "rsi_14": "RSI 14", "macd_hist_pct": "MACD Hist %",
    "dist_sma20": "vs 20DMA", "dist_sma50": "vs 50DMA", "dist_sma200": "vs 200DMA",
    "supertrend_dir": "Supertrend", "adx_14": "ADX", "di_diff": "+DI - -DI", "bb_pctb": "Bollinger %B",
    "stoch_k": "Stoch %K", "williams_r": "Williams %R", "cci_20": "CCI 20", "mfi_14": "MFI 14",
    "atr_pct": "ATR %", "vol_surge": "Volume x Avg", "obv_slope_10": "OBV Slope", "deliv_per": "Delivery %",
    "deliv_ratio": "Delivery x Avg", "turnover_cr": "Avg Turnover (Cr)", "ret_5d": "5D %", "ret_20d": "20D %",
    "ret_60d": "60D %", "rs_vs_nifty_20d": "RS vs NIFTY 20D", "dist_52w_high": "From 52W High",
    "dist_52w_low": "From 52W Low",
}
FUND_COLS = {
    "symbol": "Symbol", "company": "Company", "sector": "Sector", "market_cap_cr": "Market Cap (Cr)",
    "pe": "P/E", "forward_pe": "Fwd P/E", "pb": "P/B", "roe": "ROE", "roa": "ROA",
    "debt_to_equity": "Debt/Equity (%)", "earnings_growth": "Earnings Growth", "revenue_growth": "Revenue Growth",
    "profit_margin": "Profit Margin", "operating_margin": "Operating Margin", "dividend_yield": "Div. Yield",
    "beta": "Beta", "promoter_holding": "Insider/Promoter %", "institutional_holding": "Institutional %",
    "analyst_target": "Analyst Target", "analyst_view": "Analyst View", "analyst_count": "# Analysts",
}
PCT_HEADERS = {"1D %", "Prob. Up", "Confidence", "Expected Move %", "ML Prob. Up", "vs 20DMA", "vs 50DMA",
               "vs 200DMA", "ATR %", "5D %", "20D %", "60D %", "RS vs NIFTY 20D", "From 52W High",
               "From 52W Low", "ROE", "ROA", "Earnings Growth", "Revenue Growth", "Profit Margin",
               "Operating Margin", "Insider/Promoter %", "Institutional %", "Hit Rate", "hit_rate",
               "avg_actual_return", "Actual Return", "Max Return in Window", "Predicted Move"}


def _table(df, cols):
    present = [c for c in cols if c in df.columns]
    return df[present].rename(columns=cols)


def _format_sheet(ws, df, widths=None):
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        letter = get_column_letter(i)
        width = (widths or {}).get(col) or min(max(len(str(col)) + 2, 10), 45)
        if col in ("Key Reasons",):
            width = 110
        ws.column_dimensions[letter].width = width
        fmt = "0.0%" if col in PCT_HEADERS else ("#,##0.00" if df[col].dtype.kind == "f" else None)
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = fmt
        if col == "Prediction":
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                fill = SIGNAL_FILLS.get(row[0].value)
                if fill:
                    row[0].fill = fill
                    row[0].font = Font(bold=True)
        if col in ("Prob. Up", "Rule Score", "Expected Move %") and ws.max_row > 2:
            rng = f"{letter}2:{letter}{ws.max_row}"
            ws.conditional_formatting.add(rng, ColorScaleRule(
                start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
                mid_color="FFFFFF", end_type="max", end_color="63BE7B"))


def write_report(path, as_of, preds, market, metrics, importance, track_by_signal,
                 track_by_date, track_detail, universe_name, horizon):
    preds = preds.copy()
    if "market_cap" in preds:
        preds["market_cap_cr"] = pd.to_numeric(preds["market_cap"], errors="coerce") / 1e7
    preds = preds.sort_values("prob_up", ascending=False)
    main = _table(preds, MAIN_COLS)

    counts = preds["prediction"].value_counts()
    summary_rows = [
        ("Report", "NSE/BSE 1-2 Week Price Direction Prediction"),
        ("Data as of (market close)", str(as_of)),
        ("Generated at (IST)", str((pd.Timestamp.now("UTC").tz_localize(None) + pd.Timedelta(hours=5, minutes=30)).strftime("%Y-%m-%d %H:%M"))),
        ("Prediction horizon", f"{horizon} trading days (~{round(horizon * 7 / 5)} calendar days)"),
        ("Universe", universe_name),
        ("Stocks analysed", len(preds)),
        ("", ""),
        ("MARKET MOOD", ""),
    ] + [(k, v) for k, v in market.items()] + [
        ("", ""),
        ("SIGNAL COUNT", ""),
    ] + [(s, int(counts.get(s, 0))) for s in ["STRONG UP", "UP", "NEUTRAL", "DOWN", "STRONG DOWN"]] + [
        ("", ""),
        ("MODEL BACK-TEST (unseen recent data)", ""),
    ] + [(k, v) for k, v in metrics.items()] + [
        ("", ""),
        ("HOW TO READ", "Prob. Up = chance the price is higher after the horizon. >=62% STRONG UP, "
                        ">=55% UP, <=45% DOWN, <=38% STRONG DOWN. Target = model's expected move; "
                        "Likely Low/High = +/- ATR x sqrt(horizon)."),
        ("DISCLAIMER", "Statistical estimate for education/research only - NOT investment advice. "
                       "Markets are uncertain; even good models are wrong ~40-45% of the time. "
                       "Always use a stop loss and consult a SEBI-registered advisor."),
    ]
    summary = pd.DataFrame(summary_rows, columns=["Item", "Value"])

    sheets = {
        "Summary": summary,
        "Top Bullish": main[main["Prediction"].isin(["STRONG UP", "UP"])].head(30),
        "Top Bearish": main[main["Prediction"].isin(["STRONG DOWN", "DOWN"])].iloc[::-1].head(30),
        "All Predictions": main,
        "Technicals": _table(preds, TECH_COLS),
        "Fundamentals": _table(preds, FUND_COLS),
        "Model Importance": importance.rename(columns={"feature": "Parameter", "importance": "Importance (AUC drop)"}),
        "Track Record": track_by_signal.rename(columns={"prediction": "Prediction", "predictions": "Predictions",
                                                        "hit_rate": "Hit Rate", "avg_actual_return": "Actual Return"}),
        "Track Record by Day": track_by_date.rename(columns={"date": "Prediction Date", "predictions": "Predictions",
                                                             "hit_rate": "Hit Rate"}),
        "Past Calls Checked": track_detail,
        "Parameter Guide": pd.DataFrame(PARAMETER_GUIDE, columns=["Group", "Parameter", "What it tells you",
                                                                  "Bullish reading", "Bearish reading", "Source"]),
    }

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            if df is None or df.empty:
                df = pd.DataFrame({"Info": ["Not enough data yet - fills in after the first predictions mature."]})
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            if name == "Summary":
                ws.column_dimensions["A"].width = 42
                ws.column_dimensions["B"].width = 110
                for row in ws.iter_rows(min_row=2):
                    if row[0].value and str(row[0].value).isupper():
                        row[0].font = Font(bold=True, color="1F3864")
                    row[1].alignment = Alignment(wrap_text=True, horizontal="left")
                    if isinstance(row[1].value, float) and abs(row[1].value) <= 1.5:
                        row[1].number_format = "0.0%"
                for cell in ws[1]:
                    cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            elif name == "Parameter Guide":
                _format_sheet(ws, df, {c: 30 for c in df.columns})
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                _format_sheet(ws, df)
        ws = xw.sheets["Past Calls Checked"]
        for row in ws.iter_rows(min_row=1, max_row=1):
            for cell in row:
                if cell.value == "Correct?":
                    col = cell.column_letter
                    ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}", CellIsRule(
                        operator="equal", formula=['"YES"'], fill=SIGNAL_FILLS["UP"]))
                    ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}", CellIsRule(
                        operator="equal", formula=['"NO"'], fill=SIGNAL_FILLS["DOWN"]))
    return path
